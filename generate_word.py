
from pathlib import Path
from datetime import datetime
from collections import defaultdict
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Inches, Pt
import re

# Project files
FOLDER = Path(__file__).resolve().parent
INPUT_FILE = FOLDER / "formatted_schedule.txt"
KNOWLEDGE_BASE = FOLDER / "knowledge_base.pl"
OUTPUT_FILE = FOLDER / "Final_Schedule.docx"

# Use "ALL" or enter a professor's name.
TEACHER_NAME = "ALL"


def prolog_facts(predicate, count):
    """Read facts from the Prolog knowledge base."""
    facts = []

    if not KNOWLEDGE_BASE.exists():
        return facts

    pattern = re.compile(
        rf"^{re.escape(predicate)}\((.*)\)\.$"
    )
    atom_pattern = re.compile(r"'((?:''|[^'])*)'")

    for line in KNOWLEDGE_BASE.read_text(
        encoding="utf-8"
    ).splitlines():
        match = pattern.match(line.strip())

        if match:
            values = [
                item.replace("''", "'")
                for item in atom_pattern.findall(match.group(1))
            ]

            if len(values) == count:
                facts.append(values)

    return facts


def format_time(value):
    """Convert 24-hour time into 12-hour time."""
    value = str(value).strip()

    for pattern in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(
                value, pattern
            ).strftime("%I:%M %p").lstrip("0")
        except ValueError:
            continue

    return value


def format_office_hours(value):
    """Format faculty and syllabus office hours."""
    value = value.replace("Office Hours:", "").strip()

    # Remove day labels already stored in syllabus text.
    value = re.sub(
        r"^(Tue/Thursday|Tuesday/Thursday)\s*:\s*",
        "",
        value,
        flags=re.I
    )

    # Handle 24-hour times, such as 13:00-14:30.
    match = re.fullmatch(
        r"(\d{1,2}:\d{2})\s*-\s*(\d{1,2}:\d{2})",
        value
    )

    if match:
        start, end = match.groups()
        return f"{format_time(start)} – {format_time(end)}"

    # Handle times such as 10:45-1pm.
    match = re.fullmatch(
        r"(\d{1,2}(?::\d{2})?)\s*-\s*"
        r"(\d{1,2}(?::\d{2})?)\s*([ap]m)",
        value,
        flags=re.I
    )

    if match:
        start, end, end_period = match.groups()
        end_period = end_period.upper()

        start_hour = int(start.split(":")[0])
        end_hour = int(end.split(":")[0])

        # Infer AM for a morning-to-afternoon range.
        if end_period == "PM" and start_hour > end_hour:
            start_period = "AM"
        else:
            start_period = end_period

        if ":" not in start:
            start += ":00"

        if ":" not in end:
            end += ":00"

        return (
            f"{start} {start_period} – "
            f"{end} {end_period}"
        )

    return value


def format_days(value):
    """Convert abbreviated days into full names."""
    names = {
        "M": "Monday",
        "T": "Tuesday",
        "W": "Wednesday",
        "R": "Thursday",
        "F": "Friday"
    }

    value = value.strip().upper()

    if value in ("N/A", "UNKNOWN", ""):
        return "Online Asynchronous"

    days = [names[letter] for letter in value if letter in names]

    if not days:
        return value

    if len(days) == 1:
        return days[0]

    return ", ".join(days[:-1]) + " & " + days[-1]


def normalize_name(name):
    """Standardize professor names for matching."""
    name = re.sub(
        r"\b(dr|professor|prof)\.?\s*",
        "",
        name,
        flags=re.I
    )

    return set(re.findall(r"[a-z]+", name.lower()))


def read_courses():
    """Read the formatted class schedule."""
    courses = []
    current = None

    if not INPUT_FILE.exists():
        print("Missing formatted_schedule.txt.")
        return courses

    fields = {
        "Course ID": "course",
        "Title": "title",
        "Begin": "begin",
        "End": "end",
        "Days": "days",
        "Room": "room",
        "Instructor": "instructor"
    }

    for raw in INPUT_FILE.read_text(
        encoding="utf-8"
    ).splitlines():
        line = raw.strip()

        if line == "OFFICE HOURS":
            break

        if line.startswith("Course ID:"):
            if current:
                courses.append(current)

            current = {
                field: "" for field in fields.values()
            }

        if current and ":" in line:
            label, value = line.split(":", 1)

            if label in fields:
                current[fields[label]] = value.strip()

    if current:
        courses.append(current)

    return courses


def load_office_hours():
    """Load office hours for all professors."""
    hours = defaultdict(list)

    for professor, day, time in prolog_facts("office_hours", 3):
        entry = (day.title(), format_office_hours(time))

        if entry not in hours[professor]:
            hours[professor].append(entry)

    return hours


def add_schedule_table(document, courses):
    """Create the teaching schedule table."""
    table = document.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    headers = [
        "Course",
        "Instructor",
        "Schedule",
        "Location"
    ]

    for cell, heading in zip(table.rows[0].cells, headers):
        cell.text = heading

        for run in cell.paragraphs[0].runs:
            run.bold = True

    for course in courses:
        row = table.add_row().cells

        begin = format_time(course["begin"])
        end = format_time(course["end"])
        days = format_days(course["days"])

        if "asynchronous" in begin.lower():
            schedule = "Online Asynchronous"

        elif not end or end.lower() in ("n/a", "unknown"):
            schedule = f"{days}, {begin}"

        else:
            schedule = f"{days}, {begin} – {end}"

        row[0].text = (
            f"{course['course']}\n{course['title']}"
        )
        row[1].text = course["instructor"] or "Not provided"
        row[2].text = schedule
        row[3].text = course["room"] or "Not provided"


def add_office_hours(document, hours, selected):
    """Create the office-hours table."""
    document.add_heading("OFFICE HOURS", level=1)

    document.add_paragraph(
        "NOTE: Instructors are not to be interrupted "
        "during class times."
    )

    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"

    table.rows[0].cells[0].text = "Professor"
    table.rows[0].cells[1].text = "Office Hours"

    for cell in table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True

    found = False

    for professor in sorted(hours):
        if selected.upper() != "ALL":
            if not normalize_name(selected).issubset(
                normalize_name(professor)
            ):
                continue

        found = True

        schedule = "; ".join(
            f"{day}: {time}"
            for day, time in hours[professor]
        )

        cells = table.add_row().cells
        cells[0].text = professor
        cells[1].text = schedule

    if not found:
        cells = table.add_row().cells
        cells[0].text = selected if selected.upper() != "ALL" else "All faculty"
        cells[1].text = "Office hours not provided"


def create_document():
    """Generate the Word document."""
    courses = read_courses()
    hours = load_office_hours()

    if not courses and not hours:
        print("No schedule or office-hour information found.")
        return

    selected = TEACHER_NAME.strip()

    if selected.upper() != "ALL":
        courses = [
            course for course in courses
            if normalize_name(selected).issubset(
                normalize_name(course["instructor"])
            )
        ]

    document = Document()
    section = document.sections[0]

    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11)
    section.page_height = Inches(8.5)

    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    document.add_heading("FACULTY CLASS SCHEDULE", 0)

    document.add_paragraph(
        "Instructor: " + (
            "All Instructors"
            if selected.upper() == "ALL"
            else selected
        )
    )

    document.add_heading("TEACHING SCHEDULE", level=1)

    if courses:
        add_schedule_table(document, courses)
    else:
        document.add_paragraph("No matching classes found.")

    add_office_hours(document, hours, selected)

    try:
        document.save(OUTPUT_FILE)
        print(f"Created: {OUTPUT_FILE.name}")
        print(f"Courses included: {len(courses)}")
        print(f"Faculty with office hours: {len(hours)}")

    except PermissionError:
        print(
            "Close the Word document and run the program again."
        )


if __name__ == "__main__":
    create_document()
